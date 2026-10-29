#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gerador de documentos padronizados - Escritorio TM (Contabilidade Inteligente)
Usa o papel timbrado oficial (imagem de pagina inteira no cabecalho) como base e
aplica um padrao visual consistente para relatorios, cartas/oficios, pareceres e
layouts gerais. Gera .docx pronto; a conversao para PDF e feita a parte (soffice).
Requisitos: python-docx, e o arquivo base 'TM_base.docx' na mesma pasta.
"""

import os
from datetime import date

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1B, 0x2A, 0x57)
GOLD = RGBColor(0xC0, 0x8A, 0x1E)
GREY = RGBColor(0x55, 0x55, 0x55)
BLACK = RGBColor(0x20, 0x20, 0x20)

FONTE = "Arial"
CORPO_PT = 11
LINHA = 1.4

MESES = ["janeiro", "fevereiro", "marco", "abril", "maio", "junho",
         "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]

BASE_PADRAO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "TM_base.docx")


def data_extenso(d=None):
    d = d or date.today()
    mes = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
           "agosto", "setembro", "outubro", "novembro", "dezembro"][d.month - 1]
    return f"{d.day} de {mes} de {d.year}"


class TMDoc:
    def __init__(self, tipo="geral", base=BASE_PADRAO):
        self.tipo = tipo
        self.doc = Document(base)
        self._limpar_corpo()
        self._config_estilo_normal()

    def _limpar_corpo(self):
        body = self.doc.element.body
        for p in list(body.findall(qn("w:p"))):
            body.remove(p)

    def _config_estilo_normal(self):
        st = self.doc.styles["Normal"]
        st.font.name = FONTE
        st.font.size = Pt(CORPO_PT)
        st.font.color.rgb = BLACK
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONTE)
        pf = st.paragraph_format
        pf.space_after = Pt(0)
        pf.line_spacing = LINHA
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE

    def _novo_p(self, align=None, space_before=0, space_after=6, keep_with_next=False):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        if align is not None:
            p.alignment = align
        pf.space_before = Pt(space_before)
        pf.space_after = Pt(space_after)
        pf.line_spacing = LINHA
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        if keep_with_next:
            pf.keep_with_next = True
        return p

    def _run(self, p, texto, bold=False, italic=False, size=CORPO_PT,
             color=BLACK, font=FONTE, caps=False, spacing=None):
        r = p.add_run(texto)
        r.bold = bold
        r.italic = italic
        r.font.name = font
        r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        rpr = r._element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.insert(0, rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), font)
        if caps:
            el = OxmlElement("w:caps")
            el.set(qn("w:val"), "true")
            rpr.append(el)
        if spacing is not None:
            el = OxmlElement("w:spacing")
            el.set(qn("w:val"), str(spacing))
            rpr.append(el)
        return r

    def _borda(self, p, edge="bottom", size=6, color="C0801E", space=6):
        pPr = p._p.get_or_add_pPr()
        pbdr = pPr.find(qn("w:pBdr"))
        if pbdr is None:
            pbdr = OxmlElement("w:pBdr")
            pPr.append(pbdr)
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:space"), str(space))
        e.set(qn("w:color"), color)
        pbdr.append(e)

    def rotulo(self, texto):
        p = self._novo_p(space_after=2)
        self._run(p, texto.upper(), bold=True, size=10, color=GOLD, spacing=30)
        return self

    def titulo(self, texto):
        p = self._novo_p(space_before=0, space_after=4)
        self._run(p, texto, bold=True, size=17, color=NAVY)
        self._borda(p, "bottom", size=8, color="C0801E", space=6)
        return self

    def subtitulo(self, texto):
        p = self._novo_p(space_before=0, space_after=10)
        self._run(p, texto, italic=True, size=11.5, color=GREY)
        return self

    def local_data(self, cidade="Paiçandu/PR", quando=None):
        p = self._novo_p(align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=4, space_after=12)
        self._run(p, f"{cidade}, {data_extenso(quando)}.", size=10.5, color=GREY)
        return self

    def destinatario(self, linhas):
        for i, ln in enumerate(linhas):
            p = self._novo_p(space_after=1)
            self._run(p, ln, bold=(i == 0), size=CORPO_PT, color=BLACK)
        self.espaco(6)
        return self

    def referencia(self, texto):
        p = self._novo_p(space_before=6, space_after=10)
        self._run(p, "Ref.: ", bold=True, color=NAVY)
        self._run(p, texto, bold=True)
        return self

    def saudacao(self, texto):
        p = self._novo_p(space_before=4, space_after=10)
        self._run(p, texto)
        return self

    def secao(self, texto, nivel=1):
        p = self._novo_p(space_before=12, space_after=4, keep_with_next=True)
        if nivel == 1:
            self._run(p, texto, bold=True, size=13, color=NAVY)
        else:
            self._run(p, texto, bold=True, size=11.5, color=GOLD)
        return self

    def paragrafo(self, texto, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        p = self._novo_p(align=align, space_after=8)
        self._run(p, texto)
        return self

    def destaque(self, texto):
        p = self._novo_p(align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=4, space_after=8)
        self._run(p, texto, italic=True, color=NAVY)
        self._borda(p, "left", size=18, color="C0801E", space=10)
        p.paragraph_format.left_indent = Cm(0.4)
        return self

    def lista(self, itens):
        for it in itens:
            p = self._novo_p(space_after=3)
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            self._run(p, "•  ", bold=True, color=GOLD)
            self._run(p, it)
        return self

    def lista_numerada(self, itens):
        for i, it in enumerate(itens, 1):
            p = self._novo_p(space_after=3)
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            self._run(p, f"{i}.  ", bold=True, color=NAVY)
            self._run(p, it)
        return self

    def tabela(self, cabecalho, linhas, larguras=None):
        n = len(cabecalho)
        t = self.doc.add_table(rows=1, cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = t.rows[0].cells
        for j, c in enumerate(cabecalho):
            self._celula(hdr[j], c, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
                         fill="1B2A57", align=WD_ALIGN_PARAGRAPH.CENTER)
        for i, row in enumerate(linhas):
            cells = t.add_row().cells
            fill = "F2F4F8" if i % 2 == 0 else "FFFFFF"
            for j, val in enumerate(row):
                al = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                self._celula(cells[j], str(val), fill=fill, align=al)
        if larguras:
            for j, w in enumerate(larguras):
                for r in t.rows:
                    r.cells[j].width = Cm(w)
        self._borda_tabela(t)
        self.espaco(6)
        return self

    def _celula(self, cell, texto, bold=False, color=BLACK, fill=None,
                align=WD_ALIGN_PARAGRAPH.LEFT):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        self._run(p, texto, bold=bold, size=10.5, color=color)
        if fill:
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), fill)
            tcPr.append(shd)

    def _borda_tabela(self, t):
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "D5D9E2")
            borders.append(e)
        tblPr.append(borders)

    def fecho(self, texto="Atenciosamente,"):
        p = self._novo_p(space_before=12, space_after=4)
        self._run(p, texto)
        return self

    def assinatura(self, nome="Tony Wallichek Lopes", cargo="Contador",
                   registro="CRC PR 068568/O-8", com_linha=True, centralizado=True):
        self.espaco(24)
        align = WD_ALIGN_PARAGRAPH.CENTER if centralizado else WD_ALIGN_PARAGRAPH.LEFT
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER if centralizado else WD_TABLE_ALIGNMENT.LEFT
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "none")
            borders.append(e)
        tblPr.append(borders)
        trPr = t.rows[0]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        cell = t.rows[0].cells[0]
        cell.text = ""
        primeiro = cell.paragraphs[0]

        def _linha(texto, first=False, **kw):
            p = primeiro if first else cell.add_paragraph()
            p.alignment = align
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            self._run(p, texto, **kw)

        if com_linha:
            _linha("_" * 40, first=True, color=GREY)
            _linha(nome.upper(), bold=True, color=NAVY, size=11.5)
        else:
            _linha(nome.upper(), first=True, bold=True, color=NAVY, size=11.5)
        _linha(cargo, color=BLACK)
        if registro:
            _linha(registro, size=10, color=GREY)
        return self

    def espaco(self, pts=8):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(pts)
        p.paragraph_format.line_spacing = 1.0
        return self

    def salvar(self, caminho):
        self.doc.save(caminho)
        return caminho


if __name__ == "__main__":
    d = TMDoc(tipo="relatorio")
    d.rotulo("Documento de teste")
    d.titulo("Teste do gerador TM")
    d.local_data()
    d.paragrafo("Se voce esta lendo isto no timbrado, o gerador funcionou.")
    d.assinatura()
    d.salvar("teste_tm.docx")
    print("ok")
